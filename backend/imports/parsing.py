"""Bank statement parsing: CSV/XLSX reading, column guessing, date and amount normalization."""

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

MAX_ROWS = 5000


class ParseError(ValueError):
    pass


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ParseError("Encodage du fichier non reconnu.")


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_csv(raw: bytes):
    text = _decode(raw)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t|")
    except csv.Error:

        class dialect(csv.excel):
            delimiter = ";" if text.count(";") > text.count(",") else ","

    return [[_cell(c) for c in row] for row in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(raw: bytes):
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many different exception types
        raise ParseError("Fichier Excel illisible.") from exc
    sheet = workbook.worksheets[0]
    return [[_cell(c) for c in row] for row in sheet.iter_rows(values_only=True)]


def read_table(uploaded_file):
    """Return (headers, rows) from an uploaded .csv/.txt/.xlsx file."""
    name = (uploaded_file.name or "").lower()
    raw = uploaded_file.read()
    if name.endswith(".xlsx"):
        table = _read_xlsx(raw)
    elif name.endswith((".csv", ".txt")):
        table = _read_csv(raw)
    else:
        raise ParseError("Format non pris en charge : utilisez un fichier CSV ou Excel (.xlsx).")

    table = [row for row in table if any(cell for cell in row)]
    if len(table) < 2:
        raise ParseError("Le fichier ne contient aucune ligne de données.")

    # The header is the first row with at least two non-empty cells (skips bank banners).
    header_index = next(i for i, row in enumerate(table) if sum(1 for c in row if c) >= 2)
    headers = table[header_index]
    width = len(headers)
    rows = [(row + [""] * width)[:width] for row in table[header_index + 1 :]]
    if len(rows) > MAX_ROWS:
        raise ParseError(f"Le fichier dépasse {MAX_ROWS} lignes : découpez-le.")
    return [h or f"Colonne {i + 1}" for i, h in enumerate(headers)], rows


COLUMN_PATTERNS = {
    "date": r"date|jour|valeur|value",
    "description": r"libell|descr|label|motif|d[ée]tail|narration|op[ée]ration|intitul|r[ée]f[ée]rence",
    "debit": r"d[ée]bit|retrait|sortie|withdraw",
    "credit": r"cr[ée]dit|d[ée]p[oô]t|entr[ée]e|deposit",
    "amount": r"montant|amount|somme|valeur",
}


def guess_mapping(headers):
    """Best guess of {role: column index}; amount is only used without debit/credit."""
    mapping = {}
    lowered = [h.lower() for h in headers]
    for role, pattern in COLUMN_PATTERNS.items():
        for index, header in enumerate(lowered):
            if index in mapping.values():
                continue
            if re.search(pattern, header):
                mapping[role] = index
                break
    if "debit" in mapping and "credit" in mapping:
        mapping.pop("amount", None)
    else:
        mapping.pop("debit", None)
        mapping.pop("credit", None)
    return mapping


DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%Y/%m/%d", "%d-%m-%y")


def parse_date(value: str) -> date:
    value = value.strip().split(" ")[0].split("T")[0]
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    raise ParseError(f"Date illisible : « {value} »")


def parse_amount(value: str) -> Decimal | None:
    """Parse '1 234,56', '1.234,56', '1,234.56', '(500)', '500-', '-15 000 FCFA'. Empty → None."""
    text = value.strip()
    if not text:
        return None
    negative = text.startswith("-") or text.endswith("-") or (text.startswith("(") and text.endswith(")"))
    text = re.sub(r"[^\d,.]", "", text)
    if not text:
        return None
    if "," in text and "." in text:
        decimal_sep = "," if text.rfind(",") > text.rfind(".") else "."
        thousands = "." if decimal_sep == "," else ","
        text = text.replace(thousands, "").replace(decimal_sep, ".")
    elif "," in text:
        text = text.replace(",", ".") if re.search(r",\d{1,2}$", text) else text.replace(",", "")
    elif text.count(".") > 1 or re.search(r"\.\d{3}$", text):
        text = text.replace(".", "")
    try:
        amount = Decimal(text)
    except InvalidOperation:
        raise ParseError(f"Montant illisible : « {value} »")
    return -amount if negative else amount


def normalize_row(row, mapping):
    """Return {date, description, amount (signed)} for one raw row, or raise ParseError."""

    def cell(role):
        index = mapping.get(role)
        return row[index] if index is not None and index < len(row) else ""

    if "date" not in mapping:
        raise ParseError("Indiquez la colonne de date.")
    if "amount" in mapping:
        amount = parse_amount(cell("amount"))
    elif "debit" in mapping or "credit" in mapping:
        debit, credit = parse_amount(cell("debit")), parse_amount(cell("credit"))
        amount = (credit or Decimal("0")) - abs(debit or Decimal("0"))
    else:
        raise ParseError("Indiquez la colonne du montant (ou débit/crédit).")
    if not amount:
        raise ParseError("Montant vide ou nul.")
    return {
        "date": parse_date(cell("date")),
        "description": cell("description")[:255],
        "amount": amount,
    }
